import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_architecture_slide():
    pptx_path = r"c:\Desktop\projects\Major project\Vynk\C_B17_R1.pptx"
    prs = Presentation(pptx_path)

    # 1. Update Slide 2 (CONTENTS) to add 'System Architecture'
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame and "Abstract" in shape.text_frame.text:
            tf = shape.text_frame
            # Reconstruct contents list with System Architecture before References
            current_lines = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
            new_lines = []
            for line in current_lines:
                if line.lower() == "references" and "system architecture" not in [x.lower() for x in new_lines]:
                    new_lines.append("System Architecture")
                new_lines.append(line)
            
            # Re-populate paragraphs
            for idx, text in enumerate(new_lines):
                if idx < len(tf.paragraphs):
                    p = tf.paragraphs[idx]
                    p.text = text
                else:
                    p = tf.add_paragraph()
                    p.text = text
                p.font.name = "Calibri"
                p.font.size = Pt(20)
                p.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # 2. Add New Slide for System Architecture
    blank_layout = prs.slide_layouts[4] # Blank layout inherits master branding
    slide = prs.slides.add_slide(blank_layout)

    # Reorder to place right after Slide 13 (index 12 in 0-based), making it Slide 14
    sldIdLst = prs.slides._sldIdLst
    new_elem = sldIdLst[-1]
    sldIdLst.remove(new_elem)
    sldIdLst.insert(13, new_elem)

    # 3. Add Slide Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.28), Inches(11.33), Inches(0.55))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SYSTEM ARCHITECTURE & PROJECT FLOW"
    p.font.name = "Arial"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x0F, 0x24, 0x54) # Deep Navy
    p.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.78), Inches(11.33), Inches(0.35))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "End-to-End Pipeline: From Natural Language Intent & Multi-Attribute Profiling to Real-World Group Fulfillment"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(11)
    p_sub.font.italic = True
    p_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69) # Slate Gray
    p_sub.alignment = PP_ALIGN.CENTER

    # 4. Color Palette for 5 Architectural Tiers
    tiers = [
        {
            "header": "1. USER & INPUT LAYER",
            "bg_header": RGBColor(0x33, 0x41, 0x55), # Slate 700
            "bg_card": RGBColor(0xF8, 0xFA, 0xFC),   # Slate 50
            "border_card": RGBColor(0x94, 0xA3, 0xB8),# Slate 400
            "tag_color": RGBColor(0x47, 0x55, 0x69),
            "cards": [
                ("Natural Language Intent", "Free-form query: 'Play badminton with 4 players Saturday evening'", "NLP User Query"),
                ("Multi-Factor User Profile", "Interests, Master Skills, Weekly Availability, Locality & Budget", "Profile Attributes"),
                ("Venue & Activity Catalog", "Venue slots (timing, price) & Activity concepts / templates", "Platform Supply")
            ]
        },
        {
            "header": "2. AI INTELLIGENCE ENGINE",
            "bg_header": RGBColor(0x1D, 0x4E, 0xD8), # Blue 700
            "bg_card": RGBColor(0xEF, 0xF6, 0xFF),   # Blue 50
            "border_card": RGBColor(0x93, 0xC5, 0xFD),# Blue 300
            "tag_color": RGBColor(0x1D, 0x4E, 0xD8),
            "cards": [
                ("AI Intent Understanding", "Extracts activity taxonomy, target size & locality via LLM", "ai_service module"),
                ("Semantic Vector Embeddings", "Generates 1536-dim vector embeddings for queries & profiles", "pgvector Embedding"),
                ("AI Demand Radar", "Aggregates unmet intent signals to predict activity demand", "Demand Forecasting")
            ]
        },
        {
            "header": "3. HYBRID MATCHING CORE",
            "bg_header": RGBColor(0x0F, 0x76, 0x6E), # Teal 700
            "bg_card": RGBColor(0xF0, 0xFD, 0xFA),   # Teal 50
            "border_card": RGBColor(0x5E, 0xEB, 0xD7),# Teal 300
            "tag_color": RGBColor(0x0F, 0x76, 0x6E),
            "cards": [
                ("Deterministic Filters", "Hard-gate: Locality radius, day & time overlap, budget range", "0-100 Pts Scoring"),
                ("Cosine Semantic Matching", "High-speed cosine distance similarity across vector spaces", "Vector Matching"),
                ("Trust & Reputation Scoring", "Weights ranking by verified event completion & peer reviews", "Reputation Guard")
            ]
        },
        {
            "header": "4. GROUP & PROPOSAL LAYER",
            "bg_header": RGBColor(0x43, 0x38, 0xCA), # Indigo 700
            "bg_card": RGBColor(0xEE, 0xF2, 0xFF),   # Indigo 50
            "border_card": RGBColor(0xA5, 0xB4, 0xFC),# Indigo 300
            "tag_color": RGBColor(0x43, 0x38, 0xCA),
            "cards": [
                ("Dynamic Group Proposal", "Forms micro-community: Target Size (4) & Min Size (3)", "PROPOSED State"),
                ("Threshold Consensus Engine", "Peer invitations; members ACCEPT until minimum threshold met", "CONFIRMED State"),
                ("Controlled Action Gate", "Step 9 pipeline: Pure Validation → Permission Gate → Dispatcher", "Safe Dispatcher")
            ]
        },
        {
            "header": "5. EXECUTION & VENUE ENGINE",
            "bg_header": RGBColor(0x04, 0x78, 0x57), # Emerald 700
            "bg_card": RGBColor(0xEC, 0xFD, 0xF5),   # Emerald 50
            "border_card": RGBColor(0x6E, 0xE7, 0xB7),# Emerald 300
            "tag_color": RGBColor(0x04, 0x78, 0x57),
            "cards": [
                ("Scheduled Event Conversion", "Atomically converts confirmed proposal to official Event", "PUBLISHED Event"),
                ("Split-Cost Venue Booking", "Reserves VenueSlot, splits cost equally across all members", "VenueBooking Bridge"),
                ("Member Re-Engagement", "Detects inactive members & triggers controlled notifications", "Notification Engine")
            ]
        }
    ]

    # Geometry Setup
    start_x = Inches(0.55)
    start_y = Inches(1.18)
    col_width = Inches(2.18)
    col_gap = Inches(0.24)
    card_h = Inches(1.24)
    card_gap = Inches(0.12)
    header_h = Inches(0.40)

    # 5. Render 5 Columns
    for col_idx, tier in enumerate(tiers):
        col_x = start_x + col_idx * (col_width + col_gap)

        # Header Badge
        header_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_x, start_y, col_width, header_h)
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = tier["bg_header"]
        header_shape.line.color.rgb = tier["bg_header"]
        tf_h = header_shape.text_frame
        tf_h.word_wrap = True
        p_h = tf_h.paragraphs[0]
        p_h.text = tier["header"]
        p_h.font.name = "Arial"
        p_h.font.size = Pt(10.5)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p_h.alignment = PP_ALIGN.CENTER

        # 3 Cards per Column
        for card_idx, (card_title, card_desc, card_tag) in enumerate(tier["cards"]):
            card_y = start_y + header_h + Inches(0.08) + card_idx * (card_h + card_gap)
            card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_x, card_y, col_width, card_h)
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = tier["bg_card"]
            card_shape.line.color.rgb = tier["border_card"]
            card_shape.line.width = Pt(1.2)

            tf_c = card_shape.text_frame
            tf_c.word_wrap = True
            tf_c.margin_left = Inches(0.08)
            tf_c.margin_right = Inches(0.08)
            tf_c.margin_top = Inches(0.06)
            tf_c.margin_bottom = Inches(0.06)

            # Title
            p1 = tf_c.paragraphs[0]
            p1.text = card_title
            p1.font.name = "Arial"
            p1.font.size = Pt(10)
            p1.font.bold = True
            p1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark Navy/Black
            p1.alignment = PP_ALIGN.CENTER

            # Description
            p2 = tf_c.add_paragraph()
            p2.text = card_desc
            p2.font.name = "Calibri"
            p2.font.size = Pt(8.5)
            p2.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate text
            p2.alignment = PP_ALIGN.CENTER

            # Technical Tag
            p3 = tf_c.add_paragraph()
            p3.text = f"[{card_tag}]"
            p3.font.name = "Calibri"
            p3.font.size = Pt(8.0)
            p3.font.bold = True
            p3.font.color.rgb = tier["tag_color"]
            p3.alignment = PP_ALIGN.CENTER

        # Connector Arrow between columns (for columns 0 to 3)
        if col_idx < 4:
            arrow_x = col_x + col_width + Inches(0.04)
            arrow_y = start_y + Inches(2.20)
            arrow_w = col_gap - Inches(0.08)
            arrow_h = Inches(0.28)
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, arrow_y, arrow_w, arrow_h)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(0x64, 0x74, 0x8B) # Slate 500
            arrow.line.color.rgb = RGBColor(0x47, 0x55, 0x69)
            arrow.line.width = Pt(1)

    # 6. Bottom Closed-Loop Feedback Banner
    feedback_y = start_y + header_h + Inches(0.08) + 3 * (card_h + card_gap) + Inches(0.06)
    feedback_w = Inches(11.95)
    feedback_h = Inches(0.85)

    fb_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, feedback_y, feedback_w, feedback_h)
    fb_shape.fill.solid()
    fb_shape.fill.fore_color.rgb = RGBColor(0xF0, 0xFD, 0xF4) # Mint green 50
    fb_shape.line.color.rgb = RGBColor(0x05, 0x96, 0x69)     # Emerald border
    fb_shape.line.width = Pt(1.5)

    tf_fb = fb_shape.text_frame
    tf_fb.word_wrap = True
    tf_fb.margin_left = Inches(0.15)
    tf_fb.margin_right = Inches(0.15)
    tf_fb.margin_top = Inches(0.08)
    tf_fb.margin_bottom = Inches(0.06)

    p_fb1 = tf_fb.paragraphs[0]
    p_fb1.text = "CLOSED-LOOP FEEDBACK, REPUTATION & COMMUNITY LIFECYCLE RE-ENGAGEMENT"
    p_fb1.font.name = "Arial"
    p_fb1.font.size = Pt(10.5)
    p_fb1.font.bold = True
    p_fb1.font.color.rgb = RGBColor(0x06, 0x5F, 0x46) # Dark emerald
    p_fb1.alignment = PP_ALIGN.CENTER

    p_fb2 = tf_fb.add_paragraph()
    p_fb2.text = "Real-World Event Attendance  -->  Peer Review & Star Rating  -->  User Trust Score Recalibration  -->  Community Health Trend Analytics (Improving/Declining)  -->  AI Lifecycle Recommendations & Controlled Re-Engagement"
    p_fb2.font.name = "Calibri"
    p_fb2.font.size = Pt(9.0)
    p_fb2.font.bold = True
    p_fb2.font.color.rgb = RGBColor(0x04, 0x78, 0x57)
    p_fb2.alignment = PP_ALIGN.CENTER

    # Save presentation
    prs.save(pptx_path)
    print("Successfully added SYSTEM ARCHITECTURE & PROJECT FLOW slide!")
    print(f"Total slides now: {len(prs.slides)}")

if __name__ == "__main__":
    create_architecture_slide()
